import streamlit as st
from groq import Groq
import os
from pypdf import PdfReader
from docx import Document
import pandas as pd
from pptx import Presentation
import io

# 1. Configuración de página
st.set_page_config(page_title="EVANS.DA 🚀", page_icon="🤖", layout="centered")

# --- CSS PARA EL EFECTO CHATGPT Y BARRA FIJA ---
st.markdown("""
    <style>
    .main .block-container { padding-bottom: 200px; }
    .stFileUploadDropzone {
        min-height: 0px !important;
        padding: 5px !important;
        border-radius: 12px !important;
        background-color: #f0f2f6 !important;
        border: 1px solid #ddd !important;
    }
    .stFileUploadDropzone div div { display: none; }
    .stFileUploadDropzone::before {
        content: '📎 Adjuntar archivos para esta sesión';
        font-size: 13px;
        color: #555;
        text-align: center;
        display: block;
        padding: 5px;
    }
    </style>
    """, unsafe_allow_html=True)

# --- LÓGICA DE LOGIN ---
if "user_email" not in st.session_state:
    st.title("🔐 Acceso EVANS.DA")
    email = st.text_input("Ingresa tu correo para continuar:")
    if st.button("Ingresar"):
        if "@" in email:
            st.session_state.user_email = email
            st.rerun()
        else:
            st.error("Por favor ingresa un correo válido.")
    st.stop()

# --- INICIALIZACIÓN DE SESIÓN DE CHATS ---
if "chats" not in st.session_state:
    st.session_state.chats = {"Chat Inicial": []}
if "current_chat" not in st.session_state:
    st.session_state.current_chat = "Chat Inicial"

# --- SIDEBAR (PANEL LATERAL DE HISTORIAL) ---
with st.sidebar:
    st.write(f"👤 **Usuario:** {st.session_state.user_email}")
    st.title("📂 Mis Búsquedas")
    
    if st.button("➕ Nueva Búsqueda", use_container_width=True):
        nuevo_id = f"Búsqueda {len(st.session_state.chats) + 1}"
        st.session_state.chats[nuevo_id] = []
        st.session_state.current_chat = nuevo_id
        st.rerun()
    
    st.divider()
    
    # Lista de búsquedas guardadas
    for chat_name in st.session_state.chats.keys():
        if st.button(chat_name, key=chat_name, use_container_width=True):
            st.session_state.current_chat = chat_name
            st.rerun()

    st.spacer = st.empty()
    if st.button("🚪 Cerrar Sesión"):
        del st.session_state.user_email
        st.rerun()

# --- CUERPO PRINCIPAL ---
st.title("🤖 EVANS.DA")
st.caption(f"Conversación: {st.session_state.current_chat}")

client = Groq(api_key=st.secrets["GROQ_API_KEY"])

def procesar_archivo(file_name, file_content):
    texto = f"\n--- Fuente: {file_name} ---\n"
    try:
        if file_name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(file_content))
            for page in reader.pages: texto += (page.extract_text() or "") + "\n"
        elif file_name.endswith(".docx"):
            doc = Document(io.BytesIO(file_content))
            for para in doc.paragraphs: texto += para.text + "\n"
        elif file_name.endswith(".xlsx"):
            df = pd.read_excel(io.BytesIO(file_content))
            texto += df.to_string() + "\n"
        elif file_name.endswith(".pptx"):
            prs = Presentation(io.BytesIO(file_content))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"): texto += shape.text + "\n"
        return texto
    except Exception as e:
        return f"\nError en {file_name}: {e}\n"

# Renderizar historial del chat actual
for message in st.session_state.chats[st.session_state.current_chat]:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

# --- ÁREA DE ENTRADA FIJA ---
archivos_nuevos = st.file_uploader("", accept_multiple_files=True, label_visibility="collapsed")
prompt = st.chat_input("Escribe tu mensaje aquí...")

if prompt:
    # Guardar mensaje del usuario
    st.session_state.chats[st.session_state.current_chat].append({"role": "user", "content": prompt})
    with st.chat_message("user"):
        st.markdown(prompt)

    # Recopilar contexto (Carpeta data + Subidos)
    contexto_total = ""
    if os.path.exists("data"):
        for root, dirs, files in os.walk("data"):
            for archivo in files:
                with open(os.path.join(root, archivo), "rb") as f:
                    contexto_total += procesar_archivo(archivo, f.read())
    
    if archivos_nuevos:
        for f in archivos_nuevos:
            contexto_total += procesar_archivo(f.name, f.read())

    # Respuesta de la IA
    with st.chat_message("assistant"):
        with st.spinner("EVANS.DA está analizando..."):
            try:
                instrucciones = f"Eres EVANS.DA. Responde usando este contexto: {contexto_total[:15000]}"
                completion = client.chat.completions.create(
                    messages=[{"role": "system", "content": instrucciones}, {"role": "user", "content": prompt}],
                    model="llama-3.3-70b-versatile",
                )
                respuesta = completion.choices[0].message.content
                st.markdown(respuesta)
                # Guardar respuesta en el historial
                st.session_state.chats[st.session_state.current_chat].append({"role": "assistant", "content": respuesta})
                st.rerun()
                
            except Exception as e:
                st.error(f"Error: {e}")